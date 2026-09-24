"""Export writers (FR-018/FR-019).

Six contract columns: Clause ID · Section · Source Page · Requirement Text ·
Response · Checked. Exports are generated on demand from stored rows and
never cached on disk. The .docx writer uses python-docx when installed and
falls back to a minimal stdlib OOXML writer otherwise.
"""

from __future__ import annotations

import io
import zipfile
from html import escape

EXPORT_COLUMNS = ("Clause ID", "Section", "Source Page", "Requirement Text", "Response", "Checked")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


def _row_values(row) -> list[str]:
    return [
        row.clause_id,
        row.section,
        str(row.page),
        row.body,
        row.response or "",
        "Yes" if row.verified_at is not None else "No",
    ]


def _xlsx_safe(value: str) -> str:
    """Keep user/document text from becoming an active spreadsheet formula."""
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def build_xlsx(rows) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Compliance Matrix"
    ws.append(list(EXPORT_COLUMNS))
    for row in rows:
        ws.append([_xlsx_safe(value) for value in _row_values(row)])

    header_fill = PatternFill("solid", fgColor="173F5F")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[2].alignment = Alignment(horizontal="center", vertical="top")
        row[5].alignment = Alignment(horizontal="center", vertical="top")

    for col, width in zip("ABCDEF", (14, 36, 12, 62, 44, 10)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    if ws.max_row > 1:
        table = Table(displayName="ComplianceMatrix", ref=f"A1:F{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_DOC_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
{body}
<w:sectPr/></w:body></w:document>"""

_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""


def _cell(text: str, bold: bool = False) -> str:
    run = f"<w:r><w:rPr><w:b/></w:rPr><w:t xml:space=\"preserve\">{escape(text)}</w:t></w:r>" if bold else f"<w:r><w:t xml:space=\"preserve\">{escape(text)}</w:t></w:r>"
    return f"<w:tc><w:p>{run}</w:p></w:tc>"


def _table_xml(rows) -> str:
    parts = ["<w:tbl><w:tblPr><w:tblW w:w=\"0\" w:type=\"auto\"/></w:tblPr>"]
    parts.append("<w:tr>" + "".join(_cell(c, bold=True) for c in EXPORT_COLUMNS) + "</w:tr>")
    for row in rows:
        parts.append("<w:tr>" + "".join(_cell(v) for v in _row_values(row)) + "</w:tr>")
    parts.append("</w:tbl>")
    return "".join(parts)


def build_docx(rows) -> bytes:
    try:
        import docx  # python-docx (production image)
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        document = docx.Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)

        title = document.add_paragraph()
        title.paragraph_format.space_after = Pt(2)
        title_run = title.add_run("Compliance Matrix")
        title_run.bold = True
        title_run.font.name = "Arial"
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = RGBColor(23, 63, 95)

        note = document.add_paragraph("Source-cited draft - human verification required before use.")
        note.paragraph_format.space_after = Pt(8)
        note.runs[0].font.name = "Arial"
        note.runs[0].font.size = Pt(8.5)
        note.runs[0].font.color.rgb = RGBColor(79, 92, 105)

        table = document.add_table(rows=1, cols=len(EXPORT_COLUMNS))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = (0.7, 1.65, 0.65, 3.45, 2.55, 0.7)
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            grid_column.w = Inches(width)

        def format_cell(cell, *, header: bool, column: int) -> None:
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for edge in ("top", "left", "bottom", "right"):
                node = tc_mar.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    tc_mar.append(node)
                node.set(qn("w:w"), "80")
                node.set(qn("w:type"), "dxa")
            if header:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "173F5F")
                tc_pr.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if column in (0, 2, 5):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8 if not header else 8.5)
                    run.bold = header
                    if header:
                        run.font.color.rgb = RGBColor(255, 255, 255)

        for i, col in enumerate(EXPORT_COLUMNS):
            table.rows[0].cells[i].text = col
            format_cell(table.rows[0].cells[i], header=True, column=i)
        header_props = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_props.append(repeat)

        for row in rows:
            cells = table.add_row().cells
            for i, value in enumerate(_row_values(row)):
                cells[i].text = value
                format_cell(cells[i], header=False, column=i)
            row_props = table.rows[-1]._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            row_props.append(cant_split)

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
    except ImportError:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", _CONTENT_TYPES)
            zf.writestr("_rels/.rels", _RELS)
            zf.writestr("word/document.xml", _DOC_XML.format(body=_table_xml(rows)))
        return buffer.getvalue()
