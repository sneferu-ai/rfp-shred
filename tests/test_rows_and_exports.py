"""FR-016 row checking, FR-017 export lock, FR-018/019 exports, FR-037
corrections (AC-011/012/043)."""

import io
import json
import zipfile
from datetime import datetime, timezone

from app.core.models import Account, Correction, Requirement
from tests.conftest import (
    add_requirements,
    csrf_headers,
    make_account,
    make_rfp,
    refresh_db,
    signup_json,
)


def _entitled_matrix(db, email, rows=5):
    acct = db.query(Account).filter(Account.email == email).one()
    rfp = make_rfp(db, acct)
    rfp.unlocked_at = datetime.now(timezone.utc)
    reqs = add_requirements(db, rfp, rows)
    db.commit()
    return rfp, reqs


def test_patch_404_ordering(client, db):
    """OBL-6: exists -> ownership -> entitlement, identical 404 wording."""
    email, _ = signup_json(client)
    other = make_account(db)
    foreign_rfp = make_rfp(db, other)
    foreign_rfp.unlocked_at = datetime.now(timezone.utc)
    (foreign_row,) = add_requirements(db, foreign_rfp, 1)
    locked_rfp, locked_rows = _entitled_matrix(db, email, rows=15)
    locked_rfp.unlocked_at = None
    db.commit()
    headers = csrf_headers(client)

    # (1) nonexistent
    import uuid as uuid_mod

    resp = client.patch(f"/api/rows/{uuid_mod.uuid4()}", json={"verified": True}, headers=headers)
    assert resp.status_code == 404 and resp.json()["detail"] == "Not found"
    # (2) belongs to another account
    resp = client.patch(f"/api/rows/{foreign_row.id}", json={"verified": True}, headers=headers)
    assert resp.status_code == 404 and resp.json()["detail"] == "Not found"
    # (3) matrix not entitled
    resp = client.patch(f"/api/rows/{locked_rows[0].id}", json={"verified": True}, headers=headers)
    assert resp.status_code == 404 and resp.json()["detail"] == "Not found"


def test_patch_toggle_response_tags_and_corrections(client, db):
    email, _ = signup_json(client)
    rfp, rows = _entitled_matrix(db, email)
    row = rows[0]
    headers = csrf_headers(client)

    resp = client.patch(f"/api/rows/{row.id}", json={"verified": True, "response": "Compliant — see vol 2", "tags": ["technical", "key"]}, headers=headers)
    assert resp.status_code == 200
    refresh_db(db)
    assert row.verified_at is not None
    assert row.response == "Compliant — see vol 2"
    assert row.tags == ["technical", "key"]

    # uncheck + edit + tag change append corrections (FR-037)
    client.patch(f"/api/rows/{row.id}", json={"verified": False}, headers=headers)
    client.patch(f"/api/rows/{row.id}", json={"response": "updated"}, headers=headers)
    client.patch(f"/api/rows/{row.id}", json={"tags": ["technical"]}, headers=headers)
    refresh_db(db)
    types = [c.correction_type for c in db.query(Correction).all()]
    assert "uncheck" in types
    assert "response_edit" in types
    assert "tag_edit" in types


def test_export_lock_counts_unchecked(client, db):
    """AC-011: 409 names the unchecked count."""
    email, _ = signup_json(client)
    rfp, rows = _entitled_matrix(db, email, rows=5)
    headers = csrf_headers(client)
    for row in rows[:2]:
        client.patch(f"/api/rows/{row.id}", json={"verified": True}, headers=headers)
    resp = client.get(f"/app/matrices/{rfp.id}/export.xlsx")
    assert resp.status_code == 409
    assert "3 rows unchecked" in resp.json()["detail"]
    resp = client.get(f"/app/matrices/{rfp.id}/export.docx")
    assert resp.status_code == 409


def test_export_requires_entitlement(client, db):
    email, _ = signup_json(client)
    acct = db.query(Account).filter(Account.email == email).one()
    rfp = make_rfp(db, acct)
    add_requirements(db, rfp, 20)
    db.commit()
    resp = client.get(f"/app/matrices/{rfp.id}/export.xlsx")
    assert resp.status_code == 403


def test_xlsx_export_content(client, db):
    """AC-012 (xlsx): header row + content matches DB."""
    import openpyxl

    email, _ = signup_json(client)
    rfp, rows = _entitled_matrix(db, email, rows=4)
    headers = csrf_headers(client)
    for row in rows:
        client.patch(f"/api/rows/{row.id}", json={"verified": True, "response": "ok"}, headers=headers)
    resp = client.get(f"/app/matrices/{rfp.id}/export.xlsx")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == ["Clause ID", "Section", "Source Page", "Requirement Text", "Response", "Checked"]
    assert ws.max_row == 5
    first = [ws.cell(row=2, column=i).value for i in range(1, 7)]
    assert first[0] == rows[0].clause_id
    assert first[3] == rows[0].body
    assert first[5] == "Yes"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:F5"
    assert ws.column_dimensions["B"].width >= 30
    assert ws.column_dimensions["D"].width >= 50
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.fgColor.rgb.endswith("173F5F")
    assert ws["D2"].alignment.wrap_text is True
    assert "ComplianceMatrix" in ws.tables


def test_xlsx_export_neutralizes_formula_cells(client, db):
    """Untrusted solicitation/response text must remain inert in Excel."""
    import openpyxl

    email, _ = signup_json(client)
    rfp, rows = _entitled_matrix(db, email, rows=1)
    rows[0].clause_id = "=HYPERLINK(\"https://example.invalid\")"
    rows[0].body = "+SUM(1,1)"
    db.commit()
    headers = csrf_headers(client)
    client.patch(
        f"/api/rows/{rows[0].id}",
        json={"verified": True, "response": "@SUM(A1:A2)"},
        headers=headers,
    )

    resp = client.get(f"/app/matrices/{rfp.id}/export.xlsx")
    assert resp.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=False).active
    assert ws["A2"].value.startswith("'=")
    assert ws["D2"].value.startswith("'+")
    assert ws["E2"].value.startswith("'@")
    assert all(ws.cell(2, col).data_type != "f" for col in (1, 4, 5))


def test_docx_export_content(client, db):
    email, _ = signup_json(client)
    rfp, rows = _entitled_matrix(db, email, rows=3)
    headers = csrf_headers(client)
    for row in rows:
        client.patch(f"/api/rows/{row.id}", json={"verified": True}, headers=headers)
    resp = client.get(f"/app/matrices/{rfp.id}/export.docx")
    assert resp.status_code == 200
    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        xml = zf.read("word/document.xml").decode()
    for col in ("Clause ID", "Section", "Source Page", "Requirement Text", "Response", "Checked"):
        assert col in xml
    assert rows[0].clause_id in xml
    assert rows[1].body in xml
    assert "Compliance Matrix" in xml
    assert 'w:orient="landscape"' in xml
    assert "w:tblHeader" in xml


def test_small_matrix_exports_without_payment(client, db):
    """FR-015 <=10 clause: a 5-row matrix is entitled with no payment and
    exports after full checking (AC-026b)."""
    email, _ = signup_json(client)
    acct = db.query(Account).filter(Account.email == email).one()
    rfp = make_rfp(db, acct)
    rows = add_requirements(db, rfp, 5)
    db.commit()
    headers = csrf_headers(client)
    for row in rows:
        client.patch(f"/api/rows/{row.id}", json={"verified": True}, headers=headers)
    resp = client.get(f"/app/matrices/{rfp.id}/export.xlsx")
    assert resp.status_code == 200
